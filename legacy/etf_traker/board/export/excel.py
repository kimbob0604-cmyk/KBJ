#!/usr/bin/env python3
"""
엑셀 출력 — rankings.json 하나를 .xlsx 한 권으로 옮긴다.

입력은 `docs/RANKINGS-CONTRACT.md` 가 전부다. DB 도 state 디렉터리도 모르고,
`from_state()` 만 파일 경로를 안다 (CLAUDE.md 3장, render.py 와 같은 구조).

설계 판단 셋:

1. **숫자 칸에는 `cells_raw` 의 float 를 넣고 표시는 `number_format` 에 맡긴다.**
   계약이 표시용 문자열(`cells`)을 같이 주지만 그걸 셀에 넣으면 엑셀에서 정렬도
   조건부 서식도 안 먹는다. 사용자가 표를 받아서 다시 정렬해 보는 것이 이 산출물의
   용도라 문자열 쪽은 버린다.
2. **`null` 은 빈 칸으로 둔다.** 0 으로 채우지 않는다 (CLAUDE.md 2장 1번).
   0 으로 채우면 색 스케일에서 "계산 안 됨"이 "변동 없음"으로 둔갑한다.
3. **색 스케일의 가운데를 값이 아니라 0(비율은 100)에 고정한다.** openpyxl 기본인
   percentile 중앙값을 쓰면 전 섹터가 오른 날 표가 통째로 녹색이 되어 순위 정보가
   사라진다.

openpyxl 임포트는 모듈 최상단이 아니라 `_xl()` 안에서 한다. 이 모듈을 임포트하는
다른 진입점이 엑셀을 쓰지 않으면서도 openpyxl 없이 돌아야 하기 때문이다.
"""
import os
import re
from types import SimpleNamespace

from ..engine.config import ROOT

# ─────────────────────────── 서식 상수 ───────────────────────────
HEADER_BG = 'FFF2F2F2'          # 헤더 행 배경 (연회색)
CROSS_BG = 'FFBDD7EE'           # cross=true 종목명 칸 배경 (연파랑)
RED = 'FFF8696B'
WHITE = 'FFFFFFFF'
GREEN = 'FF63BE7B'

# 색 방향은 config/settings.yaml 의 rankings.excel_scale 이 정한다.
#   green_up  첨부 국장 시트 재현 — 상승 녹색 / 하락 적색
#   red_up    국내 관례이자 웹 대시보드와 같음 — 상승 적색 / 하락 청색
# 두 시트가 서로 어긋나 있어서 코드에 박지 않고 설정으로 뺐다. DECISIONS.md D-015.
SCALES = {
    'green_up': (RED, WHITE, GREEN),
    'red_up': ('FF6D9EEB', WHITE, RED),
}
SCALE_MID = WHITE

# kind 별 표시 규칙. 계약의 kind 표(RANKINGS-CONTRACT.md)와 1:1 이다.
NUM_FMT = {'pct': '0.00"%"', 'ratio': '0"%"', 'amount': '#,##0', 'int': '#,##0'}
# 색 스케일을 거는 kind 와 그 중앙값. 여기 없는 kind 는 색 없음.
SCALE_PIVOT = {'pct': 0.0, 'ratio': 100.0}

W_RANK = 6                      # 순위 칸
W_NUM = 10                      # 숫자 칸
W_NAME = 14                     # 종목명 칸
W_SECTOR = 18                   # 섹터명 칸
W_META = 48                     # 메타 시트 값 칸 — missing 사유가 길다

TOP_N = 5                       # 섹터별 상위 종목 칸 수. 모자라면 빈 칸으로 남는다
SHEET_MAX = 31                  # 엑셀 시트명 길이 제한
FREEZE = 'C2'                   # 헤더 1행 + 앞 2열 고정
META_SHEET = '메타'

_BAD_SHEET = re.compile(r'[\[\]:*?/\\]')


# ─────────────────────────── 지연 임포트 ───────────────────────────
def _xl():
    """openpyxl 과 이 모듈이 쓰는 스타일 객체를 한 번에 만들어 넘긴다."""
    try:
        from openpyxl import Workbook
        from openpyxl.formatting.rule import ColorScaleRule
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as ex:
        raise RuntimeError('엑셀 출력에는 openpyxl 이 필요합니다. '
                           '`pip install openpyxl` 로 설치한 뒤 다시 실행하세요.') from ex
    return SimpleNamespace(
        Workbook=Workbook,
        ColorScaleRule=ColorScaleRule,
        col=get_column_letter,
        bold=Font(bold=True),
        head_fill=PatternFill('solid', fgColor=HEADER_BG),
        cross_fill=PatternFill('solid', fgColor=CROSS_BG),
        center=Alignment(horizontal='center', vertical='center'))


# ─────────────────────────── 시트 공통 ───────────────────────────
def _wide(s):
    """한글은 화면 폭이 두 배다. 그대로 세면 헤더가 잘린다."""
    return sum(2 if ord(c) > 0x2E80 else 1 for c in str(s or ''))


def _sheet_name(wb, raw):
    """엑셀 시트명 제약(31자·금칙문자·중복)을 여기서 흡수한다.

    계약의 label/title 을 그대로 쓰다가 워크북 생성이 통째로 실패하는 쪽이 더 나쁘다.
    """
    base = _BAD_SHEET.sub(' ', str(raw or 'sheet')).strip()[:SHEET_MAX] or 'sheet'
    name, i = base, 2
    while name in wb.sheetnames:
        suf = f'-{i}'
        name = base[:SHEET_MAX - len(suf)] + suf
        i += 1
    return name


def _head(ws, xl, heads, widths):
    """헤더 행과 열 너비. 필터·틀고정은 데이터가 다 들어간 뒤 _finish 가 건다."""
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = xl.bold
        c.fill = xl.head_fill
        c.alignment = xl.center
        ws.column_dimensions[xl.col(i)].width = max(w, _wide(h) + 2)


def _finish(ws):
    """자동 필터와 틀 고정. 범위를 알아야 해서 마지막에 건다."""
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = FREEZE


def _num(ws, row, col, value, kind):
    """숫자 칸. 값이 없으면 셀을 아예 만들지 않는다 (0 으로 채우지 않는다)."""
    if value is None or isinstance(value, bool) or not isinstance(value, (int, float)):
        return
    c = ws.cell(row=row, column=col, value=float(value))
    c.number_format = NUM_FMT.get(kind, NUM_FMT['amount'])


def _scale(ws, xl, col, kind, last_row, colors=None):
    """kind 에 맞는 3색 스케일. 중앙을 0(또는 100)에 못박는 것이 핵심이다."""
    pivot = SCALE_PIVOT.get(kind)
    if pivot is None or last_row < 2:
        return
    lo, mid, hi = colors or scale_colors()
    letter = xl.col(col)
    ws.conditional_formatting.add(
        f'{letter}2:{letter}{last_row}',
        xl.ColorScaleRule(start_type='min', start_color=lo,
                          mid_type='num', mid_value=pivot, mid_color=mid,
                          end_type='max', end_color=hi))


# ─────────────────────────── 표 1·2: 섹터 ───────────────────────────
def _sector_sheet(wb, xl, board):
    """섹터 랭킹 한 판. 등수 칸에는 종목명만 넣는다 — 사용자 원본 엑셀이 그렇다."""
    ws = wb.create_sheet(_sheet_name(wb, board.get('label') or board.get('key')))
    ret_label = board.get('ret_label') or '등락률'
    heads = ['섹터', f'{ret_label} 순위', ret_label] + [f'{i}등' for i in range(1, TOP_N + 1)]
    _head(ws, xl, heads, [W_SECTOR, W_RANK, W_NUM] + [W_NAME] * TOP_N)

    row = 1
    for s in board.get('sectors') or []:
        row += 1
        ws.cell(row=row, column=1, value=s.get('name'))
        _num(ws, row, 2, s.get('rank'), 'int')
        _num(ws, row, 3, s.get('ret_raw'), 'pct')
        # top 이 5개보다 짧은 섹터가 있다. 없는 등수는 빈 칸으로 남긴다.
        for i, t in enumerate((s.get('top') or [])[:TOP_N]):
            ws.cell(row=row, column=4 + i, value=t.get('name'))
    _scale(ws, xl, 3, 'pct', row)
    _finish(ws)
    return ws


# ─────────────────────────── 표 3: 종목 ───────────────────────────
def _stock_sheet(wb, xl, board):
    """종목 랭킹 한 판. 열 구성은 계약의 columns 를 그대로 그린다."""
    ws = wb.create_sheet(_sheet_name(wb, board.get('title') or board.get('key')))
    cols = board.get('columns') or []
    heads = ['#', '종목명'] + [c.get('label') or c.get('key') for c in cols]
    _head(ws, xl, heads, [W_RANK, W_NAME] + [W_NUM] * len(cols))

    row = 1
    for r in board.get('rows') or []:
        row += 1
        _num(ws, row, 1, r.get('rank'), 'int')
        name = ws.cell(row=row, column=2, value=r.get('name'))
        # 두 표에 다 나온 종목. 사용자가 눈으로 먼저 찾는 행이라 칠해 준다.
        if r.get('cross'):
            name.fill = xl.cross_fill
        raw = r.get('cells_raw') or {}
        for i, c in enumerate(cols):
            _num(ws, row, 3 + i, raw.get(c.get('key')), c.get('kind'))
    for i, c in enumerate(cols):
        _scale(ws, xl, 3 + i, c.get('kind'), row)
    _finish(ws)
    return ws


# ─────────────────────────── 메타 ───────────────────────────
def _meta_sheet(wb, xl, rankings):
    """무슨 데이터로 언제 만든 표인지. missing 은 한 줄씩 그대로 옮긴다.

    빠진 데이터를 표 밖에 숨기지 않는 것이 규칙이다 (CLAUDE.md 2장 6번).
    """
    ws = wb.create_sheet(_sheet_name(wb, META_SHEET))
    _head(ws, xl, ['항목', '값'], [W_NAME, W_META])

    # 가중 방식은 sector_board 마다 다를 수 있어 board 별로 적는다.
    weighting = ' · '.join(
        f'{b.get("label") or b.get("key")}: {b.get("weighting")}'
        for b in rankings.get('sector_boards') or [] if b.get('weighting'))
    rows = [('as_of', rankings.get('as_of')),
            ('market', rankings.get('market')),
            ('source', rankings.get('source')),
            ('taxonomy', rankings.get('taxonomy')),
            ('generated_at', rankings.get('generated_at')),
            ('색 스케일', _scale_label()),
            ('weighting', weighting or None)]
    rows += [('missing', m) for m in rankings.get('missing') or []]
    for i, (k, v) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    _finish(ws)
    return ws


# ─────────────────────────── 진입점 ───────────────────────────
def _scale_label(cfg=None):
    """메타 시트에 적을 색 방향. 파일만 받은 사람이 색의 뜻을 알 수 있어야 한다."""
    lo, _mid, hi = scale_colors(cfg)
    return ('상승 녹색 / 하락 적색 (국장 시트 재현)' if hi == GREEN
            else '상승 적색 / 하락 청색 (국내 관례·웹 대시보드와 동일)')


def scale_colors(cfg=None):
    """설정이 고른 3색 스케일 (하단, 중앙, 상단)."""
    if cfg is None:
        from ..engine.config import load
        cfg = load()
    key = ((cfg.get('rankings') or {}).get('excel_scale')) or 'green_up'
    return SCALES.get(key, SCALES['green_up'])


def build(rankings, out_path):
    """rankings.json 하나를 .xlsx 로. 반환 out_path."""
    xl = _xl()
    wb = xl.Workbook()
    wb.remove(wb.active)        # 기본 시트는 계약에 없다
    for b in rankings.get('sector_boards') or []:
        _sector_sheet(wb, xl, b)
    for b in rankings.get('stock_boards') or []:
        _stock_sheet(wb, xl, b)
    _meta_sheet(wb, xl, rankings)

    d = os.path.dirname(os.path.abspath(out_path))
    if d:
        os.makedirs(d, exist_ok=True)
    wb.save(out_path)
    return out_path


def from_state(asof, out_path=None):
    """state/YYYYMMDD/rankings.json 을 읽어 엑셀로."""
    from ..engine.build import read
    rk = read(asof, 'rankings.json')
    if not rk:
        raise RuntimeError(f'state/{asof} 에 rankings.json 이 없다. 엔진을 먼저 돌려라.')
    out = out_path or os.path.join(ROOT, '..', 'docs', 'board', f'rankings-{asof}.xlsx')
    return build(rk, os.path.abspath(out))
