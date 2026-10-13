"""작은 엑셀 수식 평가기 — dart-report 엑셀의 '재계산 오류 0' 시험(docs/p4_design.md §4.6·§8.1).

DR CLAUDE.md "수식을 건드렸으면 반드시 재계산 검증 — 오류 0". LibreOffice 없이 같은 확인을
하려고 DR 이 쓰는 수식 모양만 계산한다: 숫자·문자열 상수, 셀 참조(`B3`·`$B$3`), 범위(`B3:B9`),
사칙(`+ - * /`)·단항 부호·괄호, 함수 `SUM`·`SUMIF`(같음 조건)·`IFERROR`. 그 밖의 함수·연산자는
`Unsupported` — 조용히 0 으로 계산하지 않는다(새 수식 모양이 생기면 평가기를 넓히라는 신호).

엑셀 규칙(필요한 만큼)
- 빈 칸은 산술에서 0, `SUM` 은 빈 칸·문자열을 건너뛴다. 산술에 문자열(빈 문자열 포함)이 오면
  `#VALUE!`(숫자 모양 문자열은 수로 바꾼다).
- 0 으로 나누기는 `#DIV/0!`. 오류는 퍼진다. `IFERROR(x, y)` 는 x 가 오류면 y.
- 순환 참조는 `#REF!`.

`evaluate(wb)` → {시트: {좌표: 값}}(수식 칸만). `errors(results)` → 최상위 결과가 오류인 칸 —
`IFERROR` 안에서 잡힌 `#DIV/0!` 는 결과가 대체값이라 여기 나오지 않는다('IFERROR 밖 `#DIV/0!` 0').
"""

from __future__ import annotations

import io
import re
from collections.abc import Callable, Iterator
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Final

from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

__all__ = ["ErrorValue", "Unsupported", "errors", "evaluate", "formula_count", "load"]


@dataclass(frozen=True)
class ErrorValue:
    code: str  # '#DIV/0!'·'#VALUE!'·'#REF!'

    def __str__(self) -> str:
        return self.code


DIV0: Final = ErrorValue("#DIV/0!")
VALUE: Final = ErrorValue("#VALUE!")
REF: Final = ErrorValue("#REF!")


class Unsupported(ValueError):
    """평가기가 모르는 수식 모양."""


_TOKEN = re.compile(
    r"""\s*(?:
      (?P<num>\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)
    | (?P<str>"(?:[^"]|"")*")
    | (?P<range>\$?[A-Z]{1,3}\$?\d+:\$?[A-Z]{1,3}\$?\d+)
    | (?P<ref>\$?[A-Z]{1,3}\$?\d+)
    | (?P<func>[A-Z][A-Z0-9.]*)\(
    | (?P<op>[-+*/(),])
    )""",
    re.VERBOSE,
)
_CELL = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")


def _tokens(text: str) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    pos = 0
    text = text.rstrip()
    while pos < len(text):
        m = _TOKEN.match(text, pos)
        if m is None or m.end() == pos:
            raise Unsupported(f"모르는 수식 조각: {text[pos : pos + 20]!r}")
        kind = m.lastgroup
        assert kind is not None
        out.append((kind, m.group(kind)))
        pos = m.end()
    return out


@dataclass(frozen=True)
class _Range:
    cells: tuple[tuple[int, int], ...]  # (row, col)


Value = float | str | ErrorValue | None


def _num(v: Value) -> float | ErrorValue:
    if v is None:
        return 0.0
    if isinstance(v, ErrorValue):
        return v
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, int | float):
        return float(v)
    try:
        return float(v.replace(",", ""))
    except ValueError:
        return VALUE


class _Sheet:
    def __init__(self, ws: Any, results: dict[str, Value]) -> None:
        self.ws = ws
        self.results = results
        self._busy: set[str] = set()

    def cell(self, row: int, col: int) -> Value:
        coord = f"{get_column_letter(col)}{row}"
        raw = self.ws.cell(row=row, column=col).value
        if isinstance(raw, str) and raw.startswith("="):
            if coord in self.results:
                return self.results[coord]
            if coord in self._busy:
                return REF
            self._busy.add(coord)
            try:
                v = _Parser(_tokens(raw[1:]), self).expr_all()
            finally:
                self._busy.discard(coord)
            self.results[coord] = v
            return v
        if raw is None or isinstance(raw, int | float | str):
            return raw if not isinstance(raw, bool) else float(raw)
        return str(raw)


class _Parser:
    def __init__(self, toks: list[tuple[str, str]], sheet: _Sheet) -> None:
        self.toks = toks
        self.i = 0
        self.sheet = sheet

    def peek(self) -> tuple[str, str] | None:
        return self.toks[self.i] if self.i < len(self.toks) else None

    def take(self, kind: str | None = None, text: str | None = None) -> tuple[str, str]:
        tok = self.peek()
        if tok is None or (kind and tok[0] != kind) or (text and tok[1] != text):
            raise Unsupported(f"수식 구문: {kind or ''}{text or ''} 자리에 {tok!r}")
        self.i += 1
        return tok

    def expr_all(self) -> Value:
        v = self.expr()
        if self.peek() is not None:
            raise Unsupported(f"수식 끝에 남은 조각: {self.toks[self.i :]!r}")
        return v

    def expr(self) -> Value:
        v = self.term()
        while (tok := self.peek()) in (("op", "+"), ("op", "-")):
            self.i += 1
            rhs = self.term()
            if tok == ("op", "+"):
                v = _arith(v, rhs, lambda a, b: a + b)
            else:
                v = _arith(v, rhs, lambda a, b: a - b)
        return v

    def term(self) -> Value:
        v = self.factor()
        while (tok := self.peek()) in (("op", "*"), ("op", "/")):
            assert tok is not None
            self.i += 1
            rhs = self.factor()
            if tok[1] == "*":
                v = _arith(v, rhs, lambda a, b: a * b)
            else:
                v = _divide(v, rhs)
        return v

    def factor(self) -> Value:
        tok = self.peek()
        if tok in (("op", "-"), ("op", "+")):
            self.i += 1
            v = _num(self.factor())
            if isinstance(v, ErrorValue):
                return v
            return -v if tok == ("op", "-") else v
        return self.primary()

    def primary(self) -> Value:
        tok = self.take()
        kind, text = tok
        if kind == "num":
            return float(text)
        if kind == "str":
            return text[1:-1].replace('""', '"')
        if kind == "ref":
            r, c = _rc(text)
            return self.sheet.cell(r, c)
        if kind == "range":
            raise Unsupported("범위는 함수 인자에서만")
        if kind == "op" and text == "(":
            v = self.expr()
            self.take("op", ")")
            return v
        if kind == "func":
            return self.call(text)
        raise Unsupported(f"모르는 조각: {tok!r}")

    def arg(self) -> Value | _Range:
        tok = self.peek()
        if tok is not None and tok[0] == "range":
            self.i += 1
            return _range(tok[1])
        return self.expr()

    def args(self) -> list[Value | _Range]:
        out: list[Value | _Range] = []
        if self.peek() == ("op", ")"):
            self.i += 1
            return out
        while True:
            out.append(self.arg())
            tok = self.take("op")
            if tok[1] == ")":
                return out
            if tok[1] != ",":
                raise Unsupported(f"인자 구분: {tok!r}")

    def values(self, a: _Range) -> Iterator[Value]:
        for r, c in a.cells:
            yield self.sheet.cell(r, c)

    def call(self, name: str) -> Value:
        args = self.args()
        if name == "SUM":
            total = 0.0
            for a in args:
                if isinstance(a, _Range):
                    for v in self.values(a):
                        if isinstance(v, ErrorValue):
                            return v
                        if isinstance(v, int | float):
                            total += float(v)
                else:
                    n = _num(a)
                    if isinstance(n, ErrorValue):
                        return n
                    total += n
            return total
        if name == "SUMIF":
            if len(args) != 3 or not isinstance(args[0], _Range) or not isinstance(args[2], _Range):
                raise Unsupported("SUMIF(범위, 조건, 합계 범위) 꼴만")
            crit = args[1]
            if not isinstance(crit, str) or crit[:1] in "<>=":
                raise Unsupported("SUMIF 조건은 같음 문자열만")
            total = 0.0
            for test, v in zip(self.values(args[0]), self.values(args[2]), strict=True):
                if isinstance(test, str) and test.lower() == crit.lower():
                    if isinstance(v, ErrorValue):
                        return v
                    if isinstance(v, int | float):
                        total += float(v)
            return total
        if name == "IFERROR":
            if len(args) != 2 or any(isinstance(a, _Range) for a in args):
                raise Unsupported("IFERROR(값, 대체) 꼴만")
            v, alt = args
            assert not isinstance(v, _Range) and not isinstance(alt, _Range)
            return alt if isinstance(v, ErrorValue) else v
        raise Unsupported(f"모르는 함수 {name}")


def _rc(ref: str) -> tuple[int, int]:
    m = _CELL.fullmatch(ref)
    if m is None:
        raise Unsupported(f"셀 참조: {ref!r}")
    return int(m.group(2)), column_index_from_string(m.group(1))


def _range(text: str) -> _Range:
    a, b = text.split(":")
    r1, c1 = _rc(a)
    r2, c2 = _rc(b)
    return _Range(
        tuple(
            (r, c)
            for r in range(min(r1, r2), max(r1, r2) + 1)
            for c in range(min(c1, c2), max(c1, c2) + 1)
        )
    )


def _arith(a: Value, b: Value, fn: Callable[[float, float], float]) -> Value:
    x, y = _num(a), _num(b)
    if isinstance(x, ErrorValue):
        return x
    if isinstance(y, ErrorValue):
        return y
    return fn(x, y)


def _divide(a: Value, b: Value) -> Value:
    x, y = _num(a), _num(b)
    if isinstance(x, ErrorValue):
        return x
    if isinstance(y, ErrorValue):
        return y
    if y == 0:
        return DIV0
    return x / y


def load(src: str | Path | bytes | Workbook) -> Workbook:
    if isinstance(src, Workbook):
        return src
    data = src if isinstance(src, bytes) else Path(src).read_bytes()
    return load_workbook(io.BytesIO(data))


def evaluate(src: str | Path | bytes | Workbook) -> dict[str, dict[str, Value]]:
    """모든 수식 칸을 계산한다 → {시트: {좌표: 값}}."""
    wb = load(src)
    out: dict[str, dict[str, Value]] = {}
    for ws in wb.worksheets:
        results: dict[str, Value] = {}
        sheet = _Sheet(ws, results)
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    sheet.cell(cell.row, cell.column)  # type: ignore[arg-type]
        out[ws.title] = dict(sorted(results.items()))
    return out


def formula_count(src: str | Path | bytes | Workbook) -> int:
    wb = load(src)
    return sum(
        1
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    )


def errors(results: dict[str, dict[str, Value]]) -> list[str]:
    """최상위 결과가 오류인 칸 — `시트!좌표=오류`."""
    return [
        f"{sheet}!{coord}={v}"
        for sheet, cells in results.items()
        for coord, v in cells.items()
        if isinstance(v, ErrorValue)
    ]
