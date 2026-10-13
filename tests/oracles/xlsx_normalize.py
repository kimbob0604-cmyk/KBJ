"""xlsx 정규화 — 엑셀 골든 비교용(docs/p4_design.md D-P4-7·§4.6 '정규화 표').

바이트 비교는 하지 않는다(CLAUDE.md §4 — 생성 시각·zip 순서·플랫폼 float 끝자리). 대신 통합
문서를 '사람이 보는 것'과 '차트가 그리는 것'으로 풀어 dict 로 만든다.

- 통합 문서: 시트 이름 순서.
- 시트: `max_row`·`max_column`, 눈금선(`showGridLines`), 고정 틀, 병합 범위, 열 너비(소수 2자리),
  행 높이.
- 셀(값이나 스타일이 있는 칸): 값 형(`n`·`s`·`f`·`b`·`d`), 값(수식은 문자열 그대로), 글꼴(이름·
  크기·굵게·기울임·색), 채움(패턴·전경색), 테두리 네 변(스타일·색), 정렬(가로·세로·줄바꿈),
  서식 문자열.
- 차트·드로잉: openpyxl 은 읽을 때 차트를 버린다 → zip 의 `xl/charts/*.xml`·`xl/drawings/**`·
  시트 `_rels` 를 `defusedxml` 로 읽어 **이름공간 URI 로 푼 태그·정렬한 속성·앞뒤 공백 지운 글**의
  정규 문자열로(C14N 과 같은 뜻 — lxml 유무로 이름공간 접두어가 달라져도 같은 값). 시리즈 참조
  범위·`axId`·`crosses`·데이터 레이블 플래그·앵커가 여기 들어 있다(DR CLAUDE.md 의 off-by-one·
  이중축 함정).
- 제외: `docProps/*`(생성·수정 시각), A열이 '생성 시각' 인 행의 B열 값(→ `"<generated_at>"`).

비교(`diff`)는 실수만 `math.isclose(rel_tol=1e-9, abs_tol=1e-9)`, 나머지는 정확히. 결과 메시지는
(경로, legacy 값, kbj 값) — 어느 시트·좌표·차트 파일·속성인지 보인다.
"""

from __future__ import annotations

import io
import math
import zipfile
from collections.abc import Mapping
from pathlib import Path
from typing import Any, Final

import defusedxml.ElementTree as DET
from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles.colors import Color

__all__ = [
    "ABS_TOL",
    "GENERATED_AT",
    "PART_PREFIXES",
    "REL_TOL",
    "canonical_xml",
    "diff",
    "normalize",
]

REL_TOL: Final = 1e-9
ABS_TOL: Final = 1e-9
GENERATED_AT: Final = "<generated_at>"
GENERATED_LABEL: Final = "생성 시각"
# 차트·드로잉과 그것을 시트에 거는 관계 파일
PART_PREFIXES: Final = ("xl/charts/", "xl/drawings/", "xl/worksheets/_rels/")


def _read(src: str | Path | bytes) -> bytes:
    if isinstance(src, bytes):
        return src
    return Path(src).read_bytes()


def _color(c: Color | None) -> str | None:
    if c is None:
        return None
    if c.type == "rgb":
        return str(c.rgb)
    if c.type == "theme":
        return f"theme:{c.theme}:{c.tint}"
    if c.type == "indexed":
        return f"indexed:{c.indexed}"
    return f"{c.type}:{c.value}"


def _style(cell: Cell) -> dict[str, Any]:
    f = cell.font
    fill = cell.fill
    b = cell.border
    a = cell.alignment
    return {
        "font": [f.name, f.sz, bool(f.b), bool(f.i), _color(f.color)],
        "fill": [fill.fill_type, _color(fill.fgColor) if fill.fill_type else None],
        "border": [
            [s.style, _color(s.color) if s.style else None]
            for s in (b.left, b.right, b.top, b.bottom)
        ],
        "align": [a.horizontal, a.vertical, bool(a.wrap_text)],
        "fmt": cell.number_format,
    }


def _value(v: Any) -> Any:
    if isinstance(v, bool | int | float | str) or v is None:
        return v
    return str(v)  # 날짜 등 — DR 은 쓰지 않는다


def _sheet(ws: Any) -> dict[str, Any]:
    cells: dict[str, Any] = {}
    gen_rows: set[int] = set()
    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell, Cell):
                continue  # 병합 안쪽 칸
            if cell.column == 1 and cell.value == GENERATED_LABEL:
                gen_rows.add(cell.row)
            if cell.value is None and not cell.has_style:
                continue
            cells[cell.coordinate] = {"t": cell.data_type, "v": _value(cell.value), **_style(cell)}
    for r in gen_rows:
        key = f"B{r}"
        if key in cells:
            cells[key]["v"] = GENERATED_AT
    widths = {
        k: [d.min, d.max, round(float(d.width), 2)]
        for k, d in sorted(ws.column_dimensions.items())
        if d.width is not None
    }
    heights = {
        str(k): round(float(d.height), 2)
        for k, d in sorted(ws.row_dimensions.items())
        if d.height is not None
    }
    return {
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "grid": ws.sheet_view.showGridLines,
        "freeze": ws.freeze_panes,
        "merged": sorted(str(m) for m in ws.merged_cells.ranges),
        "widths": widths,
        "heights": heights,
        "cells": cells,
    }


def _canon(el: Any) -> str:
    """요소 하나 → 정규 문자열(태그는 `{uri}local`, 속성 정렬, 글 앞뒤 공백 제거)."""
    attrs = "".join(f' {k}="{v}"' for k, v in sorted(el.attrib.items()))
    text = (el.text or "").strip()
    inner = text + "".join(_canon(ch) + (ch.tail or "").strip() for ch in el)
    return f"<{el.tag}{attrs}>{inner}</{el.tag}>"


def canonical_xml(data: bytes) -> str:
    """XML 바이트 → 정규 문자열. 외부 XML 이 아니어도 defusedxml 로 읽는다(규칙 하나)."""
    root = DET.fromstring(data)
    return _canon(root)


def normalize(src: str | Path | bytes) -> dict[str, Any]:
    """xlsx(경로·바이트) → 정규화 dict(`sheets`·`sheet`·`parts`)."""
    data = _read(src)
    wb = load_workbook(io.BytesIO(data))
    out: dict[str, Any] = {
        "sheets": list(wb.sheetnames),
        "sheet": {ws.title: _sheet(ws) for ws in wb.worksheets},
    }
    parts: dict[str, str] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for name in sorted(zf.namelist()):
            if name.startswith(PART_PREFIXES) and (name.endswith((".xml", ".rels"))):
                parts[name] = canonical_xml(zf.read(name))
    out["parts"] = parts
    return out


def diff(path: str, a: Any, b: Any, out: list[str], *, limit: int = 50) -> None:
    """a(legacy)·b(kbj) 를 걸어 다른 곳을 `out` 에 적는다(실수는 허용오차)."""
    if len(out) >= limit:
        return
    if isinstance(a, bool) or isinstance(b, bool):
        if type(a) is not type(b) or a != b:
            out.append(f"{path}: legacy={a!r} kbj={b!r}")
        return
    if isinstance(a, int | float) and isinstance(b, int | float):
        if not math.isclose(float(a), float(b), rel_tol=REL_TOL, abs_tol=ABS_TOL):
            out.append(f"{path}: legacy={a!r} kbj={b!r}")
        return
    if isinstance(a, Mapping) and isinstance(b, Mapping):
        ka, kb = set(a), set(b)
        for k in sorted(ka - kb, key=str):
            out.append(f"{path}.{k}: kbj 에 없다 (legacy={a[k]!r})"[:400])
        for k in sorted(kb - ka, key=str):
            out.append(f"{path}.{k}: legacy 에 없다 (kbj={b[k]!r})"[:400])
        for k in sorted(ka & kb, key=str):
            diff(f"{path}.{k}", a[k], b[k], out, limit=limit)
        return
    if isinstance(a, list | tuple) and isinstance(b, list | tuple):
        if len(a) != len(b):
            out.append(f"{path}: 길이 legacy={len(a)} kbj={len(b)}")
        for i, (x, y) in enumerate(zip(a, b, strict=False)):
            diff(f"{path}[{i}]", x, y, out, limit=limit)
        return
    if a != b:
        if isinstance(a, str) and isinstance(b, str) and len(a) > 200:
            i = next((k for k, (x, y) in enumerate(zip(a, b, strict=False)) if x != y), None)
            i = min(len(a), len(b)) if i is None else i
            out.append(
                f"{path}: 문자 {i} 에서 다르다 legacy=…{a[i : i + 120]!r} kbj=…{b[i : i + 120]!r}"
            )
        else:
            out.append(f"{path}: legacy={a!r} kbj={b!r}"[:400])
