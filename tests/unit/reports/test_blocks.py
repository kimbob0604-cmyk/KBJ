"""화면 블록(`blocks`) = 엑셀 칸 — 같은 입력에서 만든 xlsx 를 수식 평가기로 계산해 블록 값과 맞춘다.

docs/p4_design.md §4.6 '결과 payload(JSON) 추가 — 엑셀과 같은 입력에서(화면과 엑셀 숫자가 갈라지지
않게)'. 골든 G1 8사례 + G2 2회사 전부.
"""

from __future__ import annotations

import json
import math
from typing import Any

import pytest
from openpyxl.utils import get_column_letter

import kbj.engines.financials.bridge as k_bridge
import kbj.engines.financials.costs as k_costs
import kbj.engines.financials.orders as k_orders
import kbj.engines.financials.statements as k_statements
import kbj.reports.dart_excel.excel as k_excel
from kbj.data.public.dart.client import DartClient
from kbj.reports.dart_excel.pipeline import ReportInputs, collect_inputs, render
from tests.fixtures.synthetic.p4_world import build_world
from tests.golden.dart_report.golden import (
    BUILDER_CASES,
    GENERATED_AT,
    PIPELINE_COMPANIES,
    PIPELINE_YEARS,
    Mods,
    WorldDart,
    build_kwargs,
    load_case,
    world_now,
)
from tests.oracles import formula_eval as fe

MODS = Mods(k_statements, k_costs, k_bridge, k_orders, k_excel)


@pytest.fixture(scope="module")
def mapping() -> dict[str, Any]:
    return k_costs.load_mapping()


def _inputs_from_case(name: str, mapping: dict[str, Any]) -> ReportInputs:
    kw = build_kwargs(MODS, load_case(name), mapping)
    meta = kw["meta"]
    return ReportInputs(
        stock=meta["stock_code"], corp_code=meta["corp_code"], corp_name=meta["corp_name"],
        years=[2023, 2026], fs_div="CFS", fs_used=meta["fs_div_used"], periods=[],
        quarters=kw["quarters"], annual=kw["annual"], notes=[], qcosts=kw["qcosts"],
        acosts=kw["annual_costs"], unmapped=kw["unmapped"], bridge_steps=kw["bridge_steps"],
        bridge_label=kw["bridge_label"], orders=kw["orders"], mapping=mapping,
        warnings=kw["warnings"],
    )  # fmt: skip


def _v(x: Any) -> Any:
    """엑셀 평가값 → 블록 값(빈 문자열 = None)."""
    return None if x == "" else x


def _same(a: Any, b: Any) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return math.isclose(float(a), float(b), rel_tol=1e-9, abs_tol=1e-9)


def _rows(ws: Any) -> dict[str, int]:
    out: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v not in out:
            out[v] = r
    return out


def _check(inputs: ReportInputs) -> None:
    xlsx, blocks = render(inputs, GENERATED_AT)
    json.dumps(blocks, allow_nan=False)  # NaN·inf 없음
    wb = fe.load(xlsx)
    ev = fe.evaluate(wb)

    def cell(sheet: str, coord: str) -> Any:
        if coord in ev[sheet]:
            return _v(ev[sheet][coord])
        return wb[sheet][coord].value

    # 분기실적: 행 = 분기
    for j, row in enumerate(blocks["quarterly"]["rows"]):
        r = 2 + j
        assert _same(row["revenue"], cell("분기실적", f"B{r}"))
        assert _same(row["operating_income"], cell("분기실적", f"C{r}"))
        assert _same(row["opm"], cell("분기실적", f"E{r}")), (row["label"], row["opm"])
        assert _same(row["npm"], cell("분기실적", f"F{r}"))

    for sheet, key in (("연간비용구조", "annual_costs"), ("분기비용구조", "quarterly_costs")):
        blk = blocks[key]
        ws = wb[sheet]
        rows = _rows(ws)
        if "비용합계" not in rows:
            assert not blk["columns"]
            continue
        for i, _label in enumerate(blk["columns"]):
            col = get_column_letter(2 + i)
            assert _same(blk["revenue"][i], cell(sheet, f"{col}{rows['매출액']}"))
            assert _same(blk["total"][i], cell(sheet, f"{col}{rows['비용합계']}"))
            assert _same(blk["operating_income"][i], cell(sheet, f"{col}{rows['영업이익']}"))
            assert _same(blk["opm"][i], cell(sheet, f"{col}{rows['OPM']}"))
            for b, vals in blk["buckets"].items():
                assert _same(vals[i], cell(sheet, f"{col}{rows[b]}"))
        if key == "quarterly_costs":
            for i in range(len(blk["columns"])):
                col = get_column_letter(2 + i)
                assert _same(
                    blk["contribution_margin"][i], cell(sheet, f"{col}{rows['공헌이익률']}")
                )

    ob = blocks["orders"]
    assert _same(ob["total_eok"], cell("수주현황", "A4"))
    assert _same(ob["nipa_eok"], cell("수주현황", "C4"))
    assert _same(ob["nipa_ratio"], cell("수주현황", "E4"))

    steps = blocks["bridge"]["steps"]
    for j, s in enumerate(steps):
        r = 4 + j
        assert cell("이익브릿지", f"A{r}") == s["name"]
        assert _same(s["value"], cell("이익브릿지", f"B{r}"))
        assert _same(s["base"], cell("이익브릿지", f"C{r}"))
    assert blocks["warnings"] == list(inputs.warnings)
    assert blocks["source"] == "DART" and blocks["generated_at"] == GENERATED_AT.isoformat()


@pytest.mark.parametrize("name", BUILDER_CASES)
def test_빌더_사례의_블록은_엑셀_칸과_같다(name: str, mapping: dict[str, Any]) -> None:
    _check(_inputs_from_case(name, mapping))


@pytest.mark.parametrize("stock", PIPELINE_COMPANIES)
def test_파이프라인_블록은_엑셀_칸과_같다(stock: str, mapping: dict[str, Any]) -> None:
    client = WorldDart(build_world(), world_now()).client(DartClient)
    _check(collect_inputs(client, stock, PIPELINE_YEARS, "CFS", mapping, now=world_now()))


def test_잔차와_미매핑(mapping: dict[str, Any]) -> None:
    _, blocks = render(_inputs_from_case("c6_bridge_residual", mapping), GENERATED_AT)
    assert blocks["bridge"]["residual"] == pytest.approx(85.0 + 3.2 - 1.1 - 0.005, abs=0.05)
    _, b2 = render(_inputs_from_case("c2_ofs_mixed", mapping), GENERATED_AT)
    assert [u["label"] for u in b2["unmapped"]] == ["경상연구개발비"]
