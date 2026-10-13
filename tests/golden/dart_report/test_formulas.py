"""수식 검증 — 재계산 오류 0·수식 개수 고정·역산 영업이익 손계산(docs/p4_design.md §8.1).

DR CLAUDE.md '깨뜨리면 안 되는 것' 1·4: 합계·비율은 엑셀 수식이고, 수식을 건드렸으면 재계산해 오류 0
(현재 기준 수식 197개 — 스모크 사례). LibreOffice 대신 작은
평가기(`tests/oracles/formula_eval.py`)로 모든 수식 칸을 계산한다. 평가기 자체의 시험도 여기 있다.
"""

from __future__ import annotations

import io
import math
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from tests.golden.dart_report.golden import BUILDER_CASES, Mods, load_meta, run_builder_case
from tests.oracles import formula_eval as fe

BUCKETS = ("원가", "인건비", "감가상각비", "지급수수료", "외주용역비", "유지보수비", "기타")


@pytest.mark.parametrize("name", BUILDER_CASES)
def test_수식_오류_0_이고_개수가_고정이다(
    name: str, kbj_mods: Mods, mapping: dict[str, Any]
) -> None:
    _, xlsx = run_builder_case(kbj_mods, name, mapping)
    assert fe.errors(fe.evaluate(xlsx)) == []
    assert fe.formula_count(xlsx) == load_meta()["formula_counts"][name]


def test_스모크_사례는_DR_기준_197개() -> None:
    assert load_meta()["formula_counts"]["c1_smoke"] == 197


def _num(v: object) -> float:
    """원천 칸(파란 숫자) — 빈 칸은 0(엑셀 산술과 같다)."""
    if v is None:
        return 0.0
    assert isinstance(v, int | float), v
    return float(v)


def _rows_by_label(ws: Any) -> dict[str, int]:
    out: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v not in out:
            out[v] = r
    return out


@pytest.mark.parametrize("name", BUILDER_CASES)
@pytest.mark.parametrize("sheet", ["연간비용구조", "분기비용구조"])
def test_역산_영업이익은_매출_빼기_비용합계(
    name: str, sheet: str, kbj_mods: Mods, mapping: dict[str, Any]
) -> None:
    """평가한 '비용합계'·'영업이익'·'OPM' 이 입력 칸으로 손계산한 값과 같다."""
    _, xlsx = run_builder_case(kbj_mods, name, mapping)
    wb = fe.load(xlsx)
    res = fe.evaluate(wb)[sheet]
    ws = wb[sheet]
    rows = _rows_by_label(ws)
    if "비용합계" not in rows:
        assert name == "c4_empty_blocks" and sheet == "분기비용구조"  # 주석 없음 안내만
        return
    checked = 0
    for col in range(2, ws.max_column + 1):
        letter = get_column_letter(col)
        total_f = ws.cell(row=rows["비용합계"], column=col).value
        if not (isinstance(total_f, str) and total_f.startswith("=SUM(")):
            continue
        rev = _num(ws.cell(row=rows["매출액"], column=col).value)
        costs = sum(_num(ws.cell(row=rows[b], column=col).value) for b in BUCKETS)
        total = res[f"{letter}{rows['비용합계']}"]
        op = res[f"{letter}{rows['영업이익']}"]
        opm = res[f"{letter}{rows['OPM']}"]
        assert isinstance(total, float) and math.isclose(total, costs, abs_tol=1e-9)
        assert isinstance(op, float) and math.isclose(op, rev - costs, abs_tol=1e-9)
        if rev:
            assert isinstance(opm, float) and math.isclose(opm, (rev - costs) / rev)
        else:
            assert opm == ""  # IFERROR 안의 #DIV/0! → 빈칸
        checked += 1
    assert checked >= 1


# ── 평가기 자기 시험 ──────────────────────────────────────────────────────────────


def _wb(cells: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "S"
    for coord, v in cells.items():
        ws[coord] = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_평가기_기본_연산() -> None:
    res = fe.evaluate(_wb({
        "A1": 10, "A2": 5, "A3": None, "A4": "x",
        "B1": "=SUM(A1:A4)", "B2": "=A1-A2*2", "B3": "=-(A1+A2)/3", "B4": '=IFERROR(A1/A3,"")',
        "B5": "=A1/A3", "B6": "=A1+A4", "B7": '=IFERROR(A1+A4,"v")', "B8": "=B1+B2",
        "C1": "Y", "C2": "N", "C3": "y", "D1": 1, "D2": 2, "D3": 4,
        "B9": '=SUMIF(C1:C3,"Y",D1:D3)', "B10": "=(A1-(0))/A1", "B11": "=SUM($A$1:$A$2)",
    }))["S"]  # fmt: skip
    assert res["B1"] == 15.0 and res["B2"] == 0.0 and res["B3"] == -5.0 and res["B4"] == ""
    assert res["B5"] == fe.DIV0 and res["B6"] == fe.VALUE and res["B7"] == "v"
    assert res["B8"] == 15.0 and res["B9"] == 5.0 and res["B10"] == 1.0 and res["B11"] == 15.0
    assert fe.errors({"S": res}) == ["S!B5=#DIV/0!", "S!B6=#VALUE!"]


def test_평가기는_모르는_함수를_지어내지_않는다() -> None:
    with pytest.raises(fe.Unsupported):
        fe.evaluate(_wb({"A1": "=AVERAGE(B1:B2)"}))
    with pytest.raises(fe.Unsupported):
        fe.evaluate(_wb({"A1": "=A2&A3"}))


def test_평가기_순환_참조는_REF() -> None:
    res = fe.evaluate(_wb({"A1": "=A2+1", "A2": "=A1+1"}))["S"]
    assert fe.REF in res.values()
