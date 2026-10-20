"""KB자산운용 RISE — 사이트 이전(2026-10) 뒤 kbam.co.kr JSON·엑셀 API.

옛 www.riseetf.co.kr(서버 렌더링 + jQuery POST)는 2026-10 부터 전부 302 → https://kbam.co.kr 로
넘어간다(실측 2026-10-20 — etf.collect 가 '운용사 1곳 통째로 실패 — rise' 로 매일 실패했다). 새
사이트는 Next.js 이고 화면이 JSON API 를 부른다(실측 2026-10-20, 값은 적지 않는다):

- 목록: `GET /api/products/etfs?page=N&page_size=20` → `{page_items:[{fund_cd, krx_cd, name, …}],
  page_info:{total_count, total_page, next_page, …}}`. page_size 는 20 을 넘겨도 20 이다 —
  `next_page` 가 None 이 될 때까지(최대 `MAX_PAGES` — 유한 루프).
- 구성종목: `GET /api/products/etfs/{fund_cd}/holdings[?base_dt=YYYYMMDD]` → `{base_dt,
  available_dates:[YYYYMMDD…], holding_count, total_count, items:[{rank, item_cd(ISIN), item_nm,
  ratio, shares, unit_price, vamt}]}`. **상위 30줄만** 준다(total_count 가 전체 수). 공시가 없는 날
  (주말·휴장·미래)을 주면 그 전날이 아니라 **최신 날**을 조용히 준다 — 그래서 `available_dates` 에서
  요청일 이하 최신 날을 골라 다시 부른다(기준일 = 응답 `base_dt`).
- 전체 구성종목: 같은 경로 `?download=xlsx&base_dt=…` → 엑셀(시트 하나, 머리 줄 `종목코드·종목명·
  수량(주)·보유비중(%)·평가금액(원)`). 종목코드가 ISIN 이 아니라 **6자리로 줄인 값**이라 채권
  (KR6…)·원화예금(KRD…)·설정현금(CASH…)도 주식 코드처럼 보인다 — 같은 날 JSON 상위 줄의 ISIN 이
  KR7(주식·ETF)이 아닌 코드와 현금·채권·선물 이름(`_NOT_STOCK_NAMES`)은 거른다. 상위 30줄 밖의
  비주식은 이름으로만 가린다(국내 주식형만 추적하므로 드물다 — `etf_holdings` 추적 대상).

전체 줄이 JSON 에 다 있으면(total_count ≤ 받은 줄) 엑셀을 부르지 않는다. 엑셀을 못 받으면 JSON 상위
30줄로 대신하고 행 품질을 `estimated`(일부만)로 둔다 — 조용히 '전체'인 척하지 않는다.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterable, Mapping
from dataclasses import replace
from datetime import date
from typing import Any, ClassVar, Final, Literal

from openpyxl import load_workbook

from kbj.core.quality import Quality
from kbj.core.rows import HoldingRow
from kbj.data.private.etf_issuers.base import (
    FundRef,
    IssuerError,
    IssuerHttp,
    add_holding,
    holding,
    is_kr_code,
    isin_to_code,
    parse_num,
    today_kst,
    ymd,
)

MAX_PAGES: Final = 30  # 목록 쪽 상한(20 × 30 = 600 펀드)
# 엑셀에서 주식 코드처럼 보이는 비주식 줄(실측 2026-10-20: 원화예금 010010·국고채권 03501G 등)
_NOT_STOCK_NAMES: Final = (
    "원화예금", "설정현금", "현금", "예금", "미수", "미지급", "증거금", "국고채", "채권", "통안",
    "선물",
)  # fmt: skip
_HEAD: Final = ("종목코드", "종목명")


class Rise:
    KEY: ClassVar[str] = "rise"
    NAME: ClassVar[str] = "KB RISE"
    DEPTH: ClassVar[Literal["full", "top10"]] = "full"
    HISTORY: ClassVar[bool] = True
    BASE: ClassVar[str] = "https://kbam.co.kr"
    API: ClassVar[str] = BASE + "/api/products/etfs"
    PAGE_SIZE: ClassVar[int] = 20

    def __init__(self, http: IssuerHttp, *, today: Callable[[], date] = today_kst) -> None:
        self.http = http
        self.today = today
        self.headers = {"Accept": "application/json", "Referer": self.BASE + "/find"}

    # ── 목록 ──
    def universe(self) -> list[FundRef]:
        out: list[FundRef] = []
        seen: set[str] = set()
        page: int | None = 1
        for _ in range(MAX_PAGES):  # 유한 루프
            if page is None:
                break
            d = self.http.json(
                self.KEY, self.API, params={"page": page, "page_size": self.PAGE_SIZE},
                headers=self.headers,
            )  # fmt: skip
            for it in d.get("page_items") or []:
                key = str(it.get("fund_cd") or "").strip()
                if not key or key in seen:
                    continue
                seen.add(key)
                ticker = str(it.get("krx_cd") or "").strip() or None
                out.append(FundRef(key, ticker, str(it.get("name") or "").strip() or None))
            nxt = (d.get("page_info") or {}).get("next_page")
            page = int(nxt) if nxt else None
        return out

    # ── 구성종목 ──
    def _json(self, fund_key: str, day: date | None) -> Mapping[str, Any]:
        params = {"base_dt": f"{day:%Y%m%d}"} if day is not None else None
        return self.http.json(
            self.KEY, f"{self.API}/{fund_key}/holdings", params=params, headers=self.headers
        )

    def _xlsx(self, fund_key: str, day: date) -> list[tuple[str, str, float, float, float]]:
        reply = self.http.request(
            self.KEY, "GET", f"{self.API}/{fund_key}/holdings",
            params={"download": "xlsx", "base_dt": f"{day:%Y%m%d}"},
        )  # fmt: skip
        return xlsx_rows(reply.body)

    def holdings(self, fund_key: str, day: date) -> tuple[dict[str, HoldingRow], date]:
        want = min(day, self.today())  # 사이트 최신 자료는 KST 오늘을 넘을 수 없다
        first = self._json(fund_key, None)
        real = pick_date(first.get("available_dates") or [], want)
        if real is None:
            return {}, want
        d = first if ymd(first.get("base_dt")) == real else self._json(fund_key, real)
        got = ymd(d.get("base_dt"))
        if got != real:
            raise IssuerError(
                self.KEY, 200, f"{fund_key}: 기준일 {real} 을 달라 했는데 {got} 이 왔다"
            )
        items = list(d.get("items") or [])
        top, not_stock = json_rows(self.KEY, items)
        total = int(d.get("total_count") or len(items))
        if total <= len(items) or not top:
            # 다 받았거나, 비중 상위 줄에 국내 주식이 하나도 없다(채권·해외형) — 엑셀의 줄임 코드는
            # 채권도 주식처럼 보여(실측 136340 → 029788…) 부르지 않는다
            return top, real
        try:
            lines = self._xlsx(fund_key, real)
        except (IssuerError, ValueError) as e:
            # 전체를 못 받았다 — 상위 줄만 쓰고 일부라고 적는다(조용히 '전체'인 척하지 않는다)
            del e
            return {c: replace(h, quality=Quality.ESTIMATED) for c, h in top.items()}, real
        return xlsx_holdings(self.KEY, lines, not_stock), real


def pick_date(available: Iterable[object], want: date) -> date | None:
    """공시일 목록에서 want 이하 최신 날 — 없으면 None."""
    days = [d for d in (ymd(x) for x in available) if d is not None and d <= want]
    return max(days) if days else None


def json_rows(
    issuer: str, items: Iterable[Mapping[str, Any]]
) -> tuple[dict[str, HoldingRow], set[str]]:
    """JSON 상위 줄 → (국내 주식·ETF 행, 주식이 아닌 줄의 6자리 줄임 코드).

    ISIN 이 KR7 이면 단축코드(우선주 보정 — `isin_to_code`). 그 밖(KR6 채권·KR4 선물·KRD 예금 등)은
    엑셀의 줄임 코드(ISIN 3~8번째 자리)를 거르는 데 쓴다."""
    rows: dict[str, HoldingRow] = {}
    other: set[str] = set()
    for it in items:
        isin = str(it.get("item_cd") or "").strip().upper()
        code = isin_to_code(isin)
        if code is None:
            if len(isin) == 12:
                other.add(isin[3:9])
            continue
        qty, wt, val = (parse_num(it.get(k)) for k in ("shares", "ratio", "vamt"))
        add_holding(rows, holding(issuer, code, it.get("item_nm"), qty, wt, val))
    return rows, other


def xlsx_rows(body: bytes) -> list[tuple[str, str, float, float, float]]:
    """엑셀 → (코드, 이름, 수량, 비중, 금액) 줄. 머리 줄이 다르면 ValueError(모양이 바뀌었다)."""
    wb = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        head = next(it, None)
        if head is None or tuple(str(c or "").strip() for c in head[:2]) != _HEAD:
            raise ValueError(f"RISE 엑셀 머리 줄이 다르다: {head!r}")
        out: list[tuple[str, str, float, float, float]] = []
        for r in it:
            if not r or r[0] is None:
                continue
            cells = list(r) + [None] * (5 - len(r))
            out.append((str(cells[0]).strip(), str(cells[1] or "").strip(), parse_num(cells[2]),
                        parse_num(cells[3]), parse_num(cells[4])))  # fmt: skip
        return out
    finally:
        wb.close()


def xlsx_holdings(
    issuer: str, lines: Iterable[tuple[str, str, float, float, float]], not_stock: set[str]
) -> dict[str, HoldingRow]:
    """엑셀 줄 → 국내 주식·ETF 행. 현금성 이름·JSON 이 주식이 아니라고 한 코드는 뺀다."""
    rows: dict[str, HoldingRow] = {}
    for code, name, qty, wt, val in lines:
        c = code.upper()
        if not is_kr_code(c) or c in not_stock:
            continue
        if any(k in name for k in _NOT_STOCK_NAMES):
            continue
        add_holding(rows, holding(issuer, c, name, qty, wt, val))
    return rows
